"""
Export module: Excel workbook (with embedded charts) and console summary.

Sheet layout:
  1. Model Results   – full results table, colour-coded pass/fail/recommended
  2. Best Model Chart – time-series and scatter charts for the recommended model
  3. Summary          – run statistics and recommended model details
"""

import io
import warnings
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd

warnings.filterwarnings('ignore')

# Attempt to import optional but expected dependencies
try:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    MATPLOTLIB_OK = True
except ImportError:
    MATPLOTLIB_OK = False

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


# ---------------------------------------------------------------------------
# Console summary
# ---------------------------------------------------------------------------

def print_summary(results: List[Dict], best_model: Optional[Dict]) -> None:
    """Print a brief run summary to stdout."""
    passing_count = sum(1 for r in results if r['Pass / Fail'] == 'Pass')
    high_outlier  = [r for r in results if r.get('Outlier Flags', 0) > 5]

    print()
    print("=" * 62)
    print("  RUN SUMMARY")
    print("=" * 62)
    print(f"  Total models tested            : {len(results)}")
    print(f"  Models passing all thresholds  : {passing_count}")
    print(f"  Models failing                 : {len(results) - passing_count}")

    if best_model:
        print()
        print("  Recommended Model")
        print("  -----------------")
        print(f"  Model ID         : {best_model['Model ID']}")
        print(f"  R²               : {best_model['R²']:.4f}")
        print(f"  Baseline Period  : {best_model['Baseline Period']}")
        print(f"  Interval         : {best_model['Interval']}")
        print(f"  Variables        : {best_model['Independent Variables']}")
        print(f"  Model Type       : {best_model['Model Type']}")
        print(f"  Operating Mode   : {best_model['Operating Mode Segment']}")
        print(f"  Intercept        : {best_model['Intercept']}")
        if best_model.get('Change-Point Value') not in ('', None, np.nan):
            print(f"  Change-Point     : {best_model['Change-Point Value']}")
    else:
        print()
        print("  WARNING: No models passed all thresholds.")
        print("  Suggestions:")
        print("    - Review data quality and check for outliers")
        print("    - Consider additional independent variables")
        print("    - Ensure the baseline period has sufficient variation")

    if high_outlier:
        print()
        print(f"  NOTE: {len(high_outlier)} model(s) have >5 flagged outliers.")
        print("        Review the Outlier Flags column and inspect raw data.")

    print("=" * 62)


# ---------------------------------------------------------------------------
# Chart generation
# ---------------------------------------------------------------------------

def _build_chart_png(best_model: Dict, energy_col: str) -> Optional[bytes]:
    """
    Render a two-panel matplotlib figure (time series + scatter) for the
    recommended model and return it as PNG bytes.
    """
    if not MATPLOTLIB_OK:
        return None

    fitted  = best_model.get('_fitted')
    dates   = best_model.get('_dates')
    actual  = best_model.get('_actual')

    if fitted is None or actual is None or dates is None:
        return None
    if len(fitted) != len(actual):
        return None

    try:
        fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 5))

        title = (
            f"Recommended Model: {best_model['Model ID']}  |  "
            f"R² = {best_model['R²']:.4f}  |  "
            f"Variables: {best_model['Independent Variables']}  |  "
            f"Baseline: {best_model['Baseline Period']}"
        )
        fig.suptitle(title, fontsize=10, fontweight='bold')

        # --- Panel 1: Time series ---
        ax1.plot(dates, actual,            label='Actual',    color='steelblue',
                 linewidth=1.4, marker='o', markersize=3, alpha=0.9)
        ax1.plot(dates, fitted.values,     label='Predicted', color='darkorange',
                 linewidth=1.4, linestyle='--', alpha=0.9)
        ax1.set_title('Predicted vs Actual — Time Series')
        ax1.set_xlabel('Date')
        ax1.set_ylabel(energy_col)
        ax1.legend(fontsize=9)
        ax1.xaxis.set_major_formatter(mdates.DateFormatter('%Y-%m'))
        ax1.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=4, maxticks=10))
        plt.setp(ax1.xaxis.get_majorticklabels(), rotation=40, ha='right')
        ax1.grid(True, alpha=0.3)

        # --- Panel 2: Scatter ---
        ax2.scatter(actual, fitted.values, alpha=0.65, color='steelblue',
                    edgecolors='white', linewidth=0.4, s=30)
        lim_lo = min(float(np.min(actual)), float(fitted.min())) * 0.95
        lim_hi = max(float(np.max(actual)), float(fitted.max())) * 1.05
        ax2.plot([lim_lo, lim_hi], [lim_lo, lim_hi], 'r--', linewidth=1.5,
                 label='Perfect fit (45°)')
        ax2.set_xlim(lim_lo, lim_hi)
        ax2.set_ylim(lim_lo, lim_hi)
        ax2.set_title('Predicted vs Actual — Scatter')
        ax2.set_xlabel(f'Actual {energy_col}')
        ax2.set_ylabel(f'Predicted {energy_col}')
        ax2.legend(fontsize=9)
        ax2.grid(True, alpha=0.3)

        plt.tight_layout()

        buf = io.BytesIO()
        plt.savefig(buf, format='png', dpi=150, bbox_inches='tight')
        plt.close(fig)
        buf.seek(0)
        return buf.read()

    except Exception as e:
        print(f"  Warning: Chart generation failed — {e}")
        return None


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

# Colour fills
_FILL_HEADER = PatternFill('solid', fgColor='1F4E79')  if OPENPYXL_OK else None
_FILL_PASS   = PatternFill('solid', fgColor='C6EFCE')  if OPENPYXL_OK else None
_FILL_FAIL   = PatternFill('solid', fgColor='FFC7CE')  if OPENPYXL_OK else None
_FILL_BEST   = PatternFill('solid', fgColor='FFEB9C')  if OPENPYXL_OK else None
_FONT_HDR    = Font(color='FFFFFF', bold=True)          if OPENPYXL_OK else None
_FONT_BOLD   = Font(bold=True)                          if OPENPYXL_OK else None


def _col_width(values) -> int:
    """Estimate a sensible column width based on the maximum value length."""
    max_len = max((len(str(v)) for v in values if v is not None), default=8)
    return min(max_len + 2, 40)


def _build_results_df(results: List[Dict]) -> pd.DataFrame:
    """Build the export DataFrame, stripping internal '_' keys."""
    internal = {'_fitted', '_dates', '_actual'}
    rows = [{k: v for k, v in r.items() if k not in internal} for r in results]
    return pd.DataFrame(rows)


def _write_results_sheet(ws, results_df: pd.DataFrame) -> None:
    """Write the Model Results sheet with colour-coded rows."""
    headers = list(results_df.columns)
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HDR
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    pf_idx  = headers.index('Pass / Fail')  + 1 if 'Pass / Fail'  in headers else None
    rec_idx = headers.index('Recommended')  + 1 if 'Recommended'  in headers else None

    for _, row_data in results_df.iterrows():
        ws.append(list(row_data))
        row_num  = ws.max_row
        is_best  = row_data.get('Recommended')  == 'Yes'
        is_pass  = row_data.get('Pass / Fail')  == 'Pass'
        fill     = _FILL_BEST if is_best else (_FILL_PASS if is_pass else _FILL_FAIL)
        for cell in ws[row_num]:
            cell.fill = fill

    # Auto-width columns
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = _col_width(
            [c.value for c in col_cells]
        )


def _write_summary_sheet(ws, results: List[Dict], best_model: Optional[Dict]) -> None:
    """Write the Summary sheet."""
    passing = sum(1 for r in results if r['Pass / Fail'] == 'Pass')

    ws['A1'] = 'Energy Model Regression Analyzer — Run Summary'
    ws['A1'].font = Font(bold=True, size=14)

    rows = [
        [],
        ['Total Models Tested',            len(results)],
        ['Models Passing All Thresholds',   passing],
        ['Models Failing',                  len(results) - passing],
        [],
        ['— Statistical Thresholds —',      ''],
        ['p-value threshold',               '< 0.05'],
        ['t-statistic threshold (abs)',      '> 2.0'],
        ['R² threshold',                    '> 0.75'],
    ]

    if best_model:
        rows += [
            [],
            ['— Recommended Model —', ''],
            ['Model ID',         best_model['Model ID']],
            ['R²',               best_model['R²']],
            ['Baseline Period',  best_model['Baseline Period']],
            ['Interval',         best_model['Interval']],
            ['Variables',        best_model['Independent Variables']],
            ['Model Type',       best_model['Model Type']],
            ['Operating Mode',   best_model['Operating Mode Segment']],
            ['Intercept',        best_model['Intercept']],
        ]
        if best_model.get('Change-Point Value') not in ('', None, np.nan):
            rows.append(['Change-Point Value', best_model['Change-Point Value']])
    else:
        rows += [[], ['Recommended Model', 'None — no models passed all thresholds']]

    for row in rows:
        ws.append(row)

    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 28


def export_results(
    results: List[Dict],
    best_model: Optional[Dict],
    energy_col: str,
    output_path: str,
) -> None:
    """
    Export results to an Excel workbook (or CSV fallback).

    Parameters
    ----------
    results     : list of model result dicts from run_all_models().
    best_model  : the recommended model dict, or None.
    energy_col  : name of the energy column (used in chart labels).
    output_path : destination file path (.xlsx or .csv).
    """
    results_df = _build_results_df(results)

    # ---- CSV fallback -------------------------------------------------------
    if output_path.endswith('.csv') or not OPENPYXL_OK:
        csv_path = output_path if output_path.endswith('.csv') else output_path.replace('.xlsx', '.csv')
        results_df.to_csv(csv_path, index=False)
        print(f"  Results saved to : {csv_path}")
        if not OPENPYXL_OK:
            print("  (openpyxl not installed — saved as CSV instead of Excel)")
        return

    # ---- Excel workbook -----------------------------------------------------
    wb = Workbook()

    # Sheet 1: Model Results
    ws_results = wb.active
    ws_results.title = 'Model Results'
    _write_results_sheet(ws_results, results_df)

    # Sheet 2: Best Model Chart
    if best_model is not None and MATPLOTLIB_OK:
        chart_png = _build_chart_png(best_model, energy_col)
        if chart_png:
            ws_chart = wb.create_sheet('Best Model Chart')
            ws_chart['A1'] = f"Recommended Model: {best_model['Model ID']}"
            ws_chart['A1'].font = Font(bold=True, size=13)
            ws_chart['A2'] = (
                f"R² = {best_model['R²']:.4f}  |  "
                f"Variables: {best_model['Independent Variables']}  |  "
                f"Baseline: {best_model['Baseline Period']}  |  "
                f"Mode: {best_model['Operating Mode Segment']}"
            )
            img_buf = io.BytesIO(chart_png)
            img = XLImage(img_buf)
            img.anchor = 'A4'
            ws_chart.add_image(img)

    # Sheet 3: Summary
    ws_summary = wb.create_sheet('Summary')
    _write_summary_sheet(ws_summary, results, best_model)

    wb.save(output_path)
    print(f"  Results saved to : {output_path}")
